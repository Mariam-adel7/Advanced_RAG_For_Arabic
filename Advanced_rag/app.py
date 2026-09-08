from __future__ import annotations
import streamlit as st
import uuid
from io import BytesIO
from config import DEFAULT_CONFIG
from ingestion import Embedder, VectorStore
from factory import build_advanced_rag_engine
from evaluation.evaluator import Evaluator
try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False


def build_answer_docx(question: str, result, evaluation=None) -> bytes:
    doc = Document()
    title = doc.add_heading("RAG Answer Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading("Question", level=1)
    doc.add_paragraph(question or "—")

    doc.add_heading("Answer", level=1)
    for para in str(getattr(result, "answer", "") or "").split("\n"):
        if para.strip():
            doc.add_paragraph(para.strip())

    doc.add_heading("Run Details", level=1)
    route_obj = getattr(result, "route", None)
    route_name = getattr(route_obj, "route", str(route_obj or "—"))
    details_table = doc.add_table(rows=0, cols=2)
    details_table.style = "Light Grid Accent 1"

    def add_row(table, label, value):
        cells = table.add_row().cells
        cells[0].text = str(label)
        cells[1].text = str(value)

    add_row(details_table, "Route", route_name)
    add_row( details_table,
        "Confidence",
        f"{getattr(result, 'confidence_label', '—')} "
        f"({float(getattr(result, 'confidence', 0) or 0):.3f})", )
    add_row(details_table, "Latency", f"{float(getattr(result, 'latency', 0) or 0):.2f}s")
    techniques = list(getattr(result, "techniques_used", []) or [])
    add_row(details_table, "Techniques", ", ".join(techniques) if techniques else "None")

    if evaluation is not None:
        doc.add_heading("Evaluation Metrics", level=1)
        eval_table = doc.add_table(rows=0, cols=2)
        eval_table.style = "Light Grid Accent 1"
        add_row(eval_table, "Context Relevance", evaluation.context_relevance)
        add_row(eval_table, "Faithfulness", evaluation.faithfulness)
        add_row(eval_table, "Answer Relevance", evaluation.answer_relevance)
        add_row( eval_table,
                "Correctness",
                 evaluation.correctness if evaluation.correctness is not None else "Not scored (no reference answer)", )
        eval_details = getattr(evaluation, "details", {}) or {}
        judge_reason = eval_details.get("judge_reason")
        if judge_reason:
            doc.add_paragraph(f"Judge's reasoning: {judge_reason}")

    sources = list(getattr(result, "sources", []) or [])
    if sources:
        doc.add_heading("Sources", level=1)
        for source in sources:
            if isinstance(source, dict):
                src_name = source.get("source", "Unknown")
                page = source.get("page", "Unknown")
                chunk_id = source.get("chunk_id", source.get("id", "Unknown"))
                score = source.get("score", None)
            else:
                src_name = getattr(source, "source", "Unknown")
                page = getattr(source, "page", "Unknown")
                chunk_id = getattr(source, "chunk_id", "Unknown")
                score = getattr(source, "score", None)
            line = f"{src_name} — page {page} — chunk {chunk_id}"
            if score is not None:
                try:
                    line += f" — score {float(score):.3f}"
                except (TypeError, ValueError):
                    pass
            doc.add_paragraph(line, style="List Number")

    calls = list(getattr(result, "calls", []) or [])
    if calls:
        doc.add_heading("Token & Cost Usage", level=1)
        total_input = sum(int(c.get("input_tokens", 0) or 0) for c in calls)
        total_output = sum(int(c.get("output_tokens", 0) or 0) for c in calls)
        total_cost = sum(float(c.get("cost", 0) or 0) for c in calls)
        cost_table = doc.add_table(rows=0, cols=2)
        cost_table.style = "Light Grid Accent 1"
        add_row(cost_table, "Input Tokens", total_input)
        add_row(cost_table, "Output Tokens", total_output)
        add_row(cost_table, "Total Cost", f"${total_cost:.6f}")

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def docx_download_button(question: str, result, key: str, label: str = "📥 Download as Word (.docx)", evaluation=None):
    if not DOCX_AVAILABLE:
        st.caption("Install `python-docx` (`pip install python-docx`) to enable Word export.")
        return
    st.download_button(
        label=label,
        data=build_answer_docx(question, result, evaluation=evaluation),
        file_name="rag_answer.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document", key=key, )


def render_evaluation_metrics(record):
    """Show one Evaluator.EvaluationRecord as metric tiles + judge reasoning."""
    st.subheader("📊 Evaluation Metrics")
    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Context Relevance", record.context_relevance)
    c2.metric("Faithfulness", record.faithfulness)
    c3.metric("Answer Relevance", record.answer_relevance)
    c4.metric("Correctness", record.correctness if record.correctness is not None else "—")
    if record.correctness is None:
        st.caption("Correctness not scored — provide a reference answer above to enable it.")
    judge_reason = (record.details or {}).get("judge_reason")
    if judge_reason:
        st.caption(f"Judge: {judge_reason}")

st.set_page_config(
    page_title="Enterprise RAG Assistant",
    page_icon="📚",
    layout="wide",)

st.markdown(
    """
    <style>

    .title {
        font-size: 2.4rem;
        font-weight: 700;
        margin-bottom: 0;
    }

    .subtitle {
        color: #777;
        font-size: 1rem;
        margin-bottom: 25px;
    }

    .answer-box {
        padding: 20px;
        border-radius: 12px;
        border: 1px solid #ddd;
        background-color: #fafafa;
        line-height: 1.9;
        font-size: 1.05rem;
    }

    .source-box {
        padding: 15px;
        border-radius: 10px;
        border: 1px solid #ddd;
        margin-bottom: 10px;
    }

    </style>
    """, unsafe_allow_html=True,)

@st.cache_resource
def load_engine():
    embedder = Embedder( model_name=DEFAULT_CONFIG.embedding_model)
    vector_store = VectorStore.load( DEFAULT_CONFIG.store_path )
    engine = build_advanced_rag_engine( DEFAULT_CONFIG, vector_store, embedder, )
    return engine

@st.cache_resource
def load_evaluator(use_llm_judge: bool):
    engine = load_engine()
    return Evaluator(llm=engine.llm, use_llm_judge=use_llm_judge)

st.markdown( '<div class="title">📚 Enterprise RAG Assistant</div>', unsafe_allow_html=True,)
st.markdown(
    '<div class="subtitle">'
    'Arabic Enterprise Document Question Answering System'
    '</div>',
    unsafe_allow_html=True,)


with st.sidebar:
    st.header("⚙️ RAG System")
    st.write("🔀 Query Router")
    st.write("🔎 Retrieval")
    st.write("🎯 Reranking")
    st.write("📖 Context Building")
    st.write("🤖 Answer Generation")
    st.divider()
    st.write("**Configuration**")
    st.write( f"Embedding: `{DEFAULT_CONFIG.embedding_model}`" )
    st.write( f"Retrieval: `{DEFAULT_CONFIG.retrieval_strategy}`" )
    st.write( f"Chunking: `{DEFAULT_CONFIG.chunking_strategy}`" )
    st.write( f"Top-K: `{DEFAULT_CONFIG.top_k}`" )
    st.divider()
    use_llm_judge = st.checkbox(
        "🧑‍⚖️ Use LLM judge for evaluation",
        value=True,
        help="When on, an LLM scores context relevance, faithfulness, answer "
             "relevance and correctness for every answer (adds one extra LLM "
             "call and a little latency/cost per question). When off, only "
             "the fast deterministic heuristics are used.",)

MODE_CHAT = "💬 محادثة (Chat)"
MODE_QA = "📊 سؤال وجواب تفصيلي (QA)"

mode = st.radio(
    "اختر طريقة الاستخدام / Choose mode:",
    [MODE_CHAT, MODE_QA],
    horizontal=True,
    key="app_mode",)
st.divider()

# =============================================================================
# CHAT MODE — multi-turn conversation, engine keeps chat_history across turns
# =============================================================================
if mode == MODE_CHAT:
    if "chat_messages" not in st.session_state:
        st.session_state.chat_messages = []
    if "chat_last_result" not in st.session_state:
        st.session_state.chat_last_result = {}  
    if "chat_evaluations" not in st.session_state:
        st.session_state.chat_evaluations = {}

    top_col1, top_col2 = st.columns([5, 1])
    with top_col2:
        if st.button("🗑️ مسح المحادثة", use_container_width=True):
            st.session_state.chat_messages = []
            st.session_state.chat_last_result = {}
            st.session_state.chat_evaluations = {}
            st.rerun()

    for i, msg in enumerate(st.session_state.chat_messages):
        with st.chat_message(msg["role"]):
            st.markdown(msg["content"])
            if msg["role"] == "assistant" and i in st.session_state.chat_last_result:
                result = st.session_state.chat_last_result[i]
                with st.expander("ℹ️ التفاصيل"):
                    d1, d2, d3 = st.columns(3)
                    d1.metric("Route", result.route.route)
                    d2.metric("Confidence", result.confidence_label)
                    d3.metric("Latency", f"{result.latency:.2f}s")
                    if result.techniques_used:
                        st.write("**Techniques:**", ", ".join(result.techniques_used))
                    if result.sources:
                        st.write("**Sources:**")
                        for j, source in enumerate(result.sources, start=1):
                            if isinstance(source, dict):
                                st.write(f"{j}. {source.get('source', 'Unknown')} " f"(page {source.get('page', '?')})")
                            else:
                                st.write(f"{j}. {source}")
                    if i in st.session_state.chat_evaluations:
                        render_evaluation_metrics(st.session_state.chat_evaluations[i])
                    prior_user_msg = st.session_state.chat_messages[i - 1]["content"] if i > 0 else ""
                    docx_download_button(
                        getattr(result, "question", prior_user_msg),
                        result,
                        key=f"chat_docx_{i}",
                        evaluation=st.session_state.chat_evaluations.get(i),
                    )

    user_input = st.chat_input("اكتب سؤالك هنا... / Type your question...")

    if user_input:
        prior_history = list(st.session_state.chat_messages)

        st.session_state.chat_messages.append({"role": "user", "content": user_input})
        with st.chat_message("user"):
            st.markdown(user_input)

        with st.chat_message("assistant"):
            try:
                with st.spinner("جارٍ التفكير..."):
                    engine = load_engine()
                    result = engine.answer(user_input, chat_history=prior_history)
                st.markdown(result.answer)
            except Exception as e:
                st.error("حدث خطأ أثناء تشغيل نظام الأسئلة والأجوبة.")
                st.exception(e)
                result = None

        if result is not None:
            assistant_index = len(st.session_state.chat_messages)
            st.session_state.chat_messages.append( {"role": "assistant", "content": result.answer} )
            st.session_state.chat_last_result[assistant_index] = result
            evaluator = load_evaluator(use_llm_judge)
            st.session_state.chat_evaluations[assistant_index] = evaluator.evaluate( str(uuid.uuid4()), result )
            st.rerun()

# =============================================================================
# QA MODE — single question, full diagnostic breakdown (previous behavior)
# =============================================================================
else:
    question = st.text_area(
        "Ask your question",
        placeholder=(
            "مثال: ما المقصود بالمخزون الراكد؟\n"
            "Example: What is the procedure for stagnant inventory?" ), height=110,)

    reference_answer = st.text_area(
        "Reference answer (optional)",
        placeholder="Paste a known-correct answer here to get a real Correctness score, "
                    "instead of one scored conservatively without ground truth.", height=80, )

    ask_button = st.button( "🔍 Ask", type="primary", use_container_width=True, )

    if ask_button:
        if not question.strip():
            st.warning( "Please enter a question." )
        else:
            try:
                with st.spinner( "Running the RAG pipeline..." ):
                    engine = load_engine()
                    result = engine.answer( question.strip() )
                    evaluator = load_evaluator(use_llm_judge)
                    evaluation = evaluator.evaluate( str(uuid.uuid4()), result, reference_answer=reference_answer.strip() or None )

                st.session_state["result"] = result
                st.session_state["evaluation"] = evaluation
            except Exception as e:
                st.error( "An error occurred while running the RAG pipeline." )
                st.exception(e)

    if "result" in st.session_state:
        result = st.session_state["result"]
        evaluation = st.session_state.get("evaluation")
        st.divider()
        st.subheader("🤖 Answer")
        st.markdown( f""" <div class="answer-box"> {result.answer} </div> """, unsafe_allow_html=True, )
        docx_download_button(getattr(result, "question", question), result, key="qa_docx", evaluation=evaluation)

        st.subheader("📊 RAG Results")
        col1, col2, col3, col4 = st.columns(4)
        with col1:
            st.metric( "Route", result.route.route,)
        with col2:
            st.metric( "Confidence", result.confidence_label,)
        with col3:
            st.metric( "Confidence Score", f"{result.confidence:.3f}", )
        with col4:
            st.metric( "Latency", f"{result.latency:.2f}s", )

        if evaluation is not None:
            render_evaluation_metrics(evaluation)

        st.subheader("🛠 Techniques Used")
        if result.techniques_used:
            cols = st.columns( min( len(result.techniques_used), 4, ))
            for i, technique in enumerate( result.techniques_used ):
                cols[ i % len(cols)].info( technique )
        else:
            st.info( "No advanced techniques were used." )


        st.subheader("📄 Sources")
        if result.sources:
            for i, source in enumerate(result.sources, start=1):
                with st.expander(f"Source {i}"):
                    if isinstance(source, dict):
                        st.write( "Source:", source.get("source", "Unknown"), )
                        st.write( "Page:", source.get("page", "Unknown"), )
                        st.write( "Chunk:", source.get("chunk_id",source.get("id", "Unknown"), ),)
                        st.json(source)
                    else:
                        st.write(source)

        else:
            st.info("No sources returned.")

        st.subheader("💰 Token & Cost Usage")
        total_input = 0
        total_output = 0
        total_cost = 0.0

        for call in result.calls:
            total_input += int( call.get( "input_tokens", 0, ))
            total_output += int( call.get( "output_tokens", 0,))
            total_cost += float( call.get( "cost", 0,))
        col1, col2, col3 = st.columns(3)

        with col1:
            st.metric( "Input Tokens", total_input, )
        with col2:
            st.metric( "Output Tokens", total_output, )
        with col3:
            st.metric( "Total Cost", f"${total_cost:.6f}", )


        with st.expander( "🔎 Retrieved Chunks"):
            if result.retrieved:
                for i, hit in enumerate( result.retrieved, start=1,):
                    chunk = hit.chunk
                    st.markdown( f"### Chunk {i}")
                    st.write( "**Chunk ID:**", chunk.chunk_id,)
                    st.write( "**Type:**", chunk.chunk_type,)
                    st.write( "**Score:**", round(float(hit.score),4, ),)
                    st.write( "**Page:**", getattr(chunk,"page",None, ),)
                    st.write( "**Source:**", getattr(chunk,"source",None, ), )
                    st.text( chunk.text )
                    st.divider()
            else:
                st.info( "No chunks retrieved." )


        with st.expander( "📖 Final Context" ):
            st.text( result.context )

        with st.expander( "⚙️ Advanced Pipeline Details"):
            st.json( result.intermediate )

        with st.expander( "🤖 LLM Calls" ):
            if result.calls:
                for i, call in enumerate( result.calls, start=1, ):
                    st.markdown( f"### Call {i}" )
                    st.json( call )
            else:
                st.info( "No LLM calls." )