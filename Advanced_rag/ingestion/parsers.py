from __future__ import annotations
import csv
import os
import re
from abc import ABC, abstractmethod
from typing import List, Optional, Tuple
from .schema import Document, Element, Table
import pytesseract
try:
    import cv2
    import numpy as np
except ImportError:
    cv2 = None
    np = None
pytesseract.pytesseract.tesseract_cmd = os.getenv( "TESSERACT_CMD", r"C:\Program Files\Tesseract-OCR\tesseract.exe", )

class BaseParser(ABC):
    @abstractmethod
    def parse(self, path: str) -> Document:
        ...

class TextParser(BaseParser):
    def parse(self, path: str) -> Document:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            raw = f.read()
        elements = [
            Element(type="paragraph", text=p.strip())
            for p in re.split(r"\n\s*\n", raw)
            if p.strip()
        ]
        return Document.create(source=path, elements=elements, format="txt")

class MarkdownParser(BaseParser):
    _HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
    _TABLE_ROW_RE = re.compile(r"^\|(.+)\|\s*$")
    _TABLE_SEP_RE = re.compile( r"^\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?$" )

    def parse(self, path: str) -> Document:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            lines = f.read().splitlines()

        elements: List[Element] = []
        buffer: List[str] = []
        def flush_paragraph() -> None:
            text = "\n".join(buffer).strip()
            if text:
                elements.append(Element(type="paragraph", text=text))
            buffer.clear()

        i = 0
        while i < len(lines):
            line = lines[i]
            heading_match = self._HEADING_RE.match(line)
            if heading_match:
                flush_paragraph()
                elements.append(
                    Element( type="heading", text=heading_match.group(2).strip(), level=len(heading_match.group(1)), ))
                i += 1
                continue

            if ( self._TABLE_ROW_RE.match(line) and i + 1 < len(lines) and self._TABLE_SEP_RE.match(lines[i + 1]) ):
                flush_paragraph()
                table_lines = [line]
                i += 2
                while i < len(lines) and self._TABLE_ROW_RE.match(lines[i]):
                    table_lines.append(lines[i])
                    i += 1
                rows = [ [cell.strip() for cell in row.strip().strip("|").split("|")] for row in table_lines ]
                elements.append(Element(type="table", text="", table=Table(rows=rows)))
                continue

            if not line.strip():
                flush_paragraph()
            else:
                buffer.append(line)
            i += 1
        flush_paragraph()
        return Document.create(source=path, elements=elements, format="markdown")

class HTMLParser(BaseParser):
    def parse(self, path: str) -> Document:
        from bs4 import BeautifulSoup
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            soup = BeautifulSoup(f.read(), "html.parser")
        elements: List[Element] = []
        body = soup.body or soup

        for tag in body.find_all( ["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "table"] ):
            if tag.name == "table":
                rows = []
                for tr in tag.find_all("tr"):
                    cells = tr.find_all(["td", "th"])
                    row = [cell.get_text(" ", strip=True) for cell in cells]
                    if any(row):
                        rows.append(row)
                if rows:
                    elements.append(Element(type="table", text="", table=Table(rows=rows)))
            elif tag.name.startswith("h"):
                text = tag.get_text(" ", strip=True)
                if text:
                    elements.append( Element(type="heading", text=text, level=int(tag.name[1])) )
            else:
                text = tag.get_text(" ", strip=True)
                if text:
                    elements.append( Element(type="list_item" if tag.name == "li" else "paragraph", text=text) )

        return Document.create(source=path, elements=elements, format="html")

class CSVParser(BaseParser):
    def parse(self, path: str) -> Document:
        with open(path, "r", encoding="utf-8", errors="replace", newline="") as f:
            reader = csv.reader(f)
            rows = [row for row in reader if any(cell.strip() for cell in row)]
        elements =[ Element(type="table", text="", table=Table(rows=rows, caption=os.path.basename(path))) ]
        return Document.create(source=path, elements=elements, format="csv")

class PDFParser(BaseParser):
    OCR_LANG = "ara+eng"
    OCR_CONFIG = r"--oem 3 --psm 6"
    TABLE_OCR_CONFIG = r"--oem 3 --psm 6"
    CELL_OCR_CONFIG = r"--oem 3 --psm 6"
    _PROCEDURE_HEADING_RE = re.compile(
        r"^\s*(?:[0-9٠-٩]+\s*[.)-]?\s*)?"
        r"(?:إجراءات|إجراء|الإجراءات|الإجراء|العملية|العمليات|المهمة|المهام)\b", re.IGNORECASE, )
    _NUMBERED_HEADING_RE = re.compile( r"^\s*(?:[0-9٠-٩]{1,3}\s*[.)-])\s+[^.]{2,160}$")
    _SECTION_HINT_RE = re.compile(
        r"^(?:[0-9٠-٩]+\s*)?(?:مقدمة|التعريف|الهدف|النطاق|المسؤوليات|المسؤولية|"
        r"السياسات|السياسة|الإجراءات|المراجع|النماذج|التقارير|الملاحق|جدول التعديلات)\b", re.IGNORECASE, )

    def __init__(self, dpi: int = 260):
        self.dpi = dpi

    @staticmethod
    def _normalize_ocr_text(text: str) -> str:
        if not text:
            return ""
        text = text.replace("\ufeff", "")
        text = text.replace("\u200e", "")
        text = text.replace("\u200f", "")
        text = text.replace("\xa0", " ")
        text = re.sub(r"[ \t]+", " ", text)
        text = re.sub(r"\n[ \t]*\n[ \t]*\n+", "\n\n", text)
        return text.strip()

    @staticmethod
    def _preprocess_ocr_image(image):
        if cv2 is None or np is None:
            return image
        img = np.array(image)
        if img.ndim == 3:
            gray = cv2.cvtColor(img, cv2.COLOR_RGB2GRAY)
        else:
            gray = img
        gray = cv2.medianBlur(gray, 3)
        gray = cv2.normalize(gray, None, 0, 255, cv2.NORM_MINMAX)
        _, binary = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)
        return binary

    @staticmethod
    def _remove_ocr_layout_noise(text: str) -> str:
        if not text:
            return ""
        lines = []
        previous = None
        for raw_line in text.splitlines():
            line = raw_line.strip()
            if not line:
                if lines and lines[-1] != "":
                    lines.append("")
                continue

            if not re.search(r"[\u0600-\u06FFA-Za-z0-9]", line):
                continue
            if previous == line:
                continue

            lines.append(line)
            previous = line

        return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()

    @staticmethod
    def _confidence_stats(data) -> Tuple[float, int, int]:
        values = []
        for conf, text in zip(data.get("conf", []), data.get("text", [])):
            token = (text or "").strip()
            try:
                score = float(conf)
            except (TypeError, ValueError):
                continue
            if token:
                values.append((score, token))

        if not values:
            return 0.0, 0, 0
        mean_conf = sum(score for score, _ in values) / len(values)
        low_count = sum(1 for score, token in values if score < 25 and len(token) <= 4)
        return mean_conf, low_count, len(values)

    @classmethod
    def _clean_low_confidence_tokens(cls, data) -> str:
        lines = {}
        n = len(data.get("text", []))

        for i in range(n):
            token = (data.get("text", [""] * n)[i] or "").strip()
            if not token:
                continue
            try:
                conf = float(data.get("conf", ["-1"] * n)[i])
            except (TypeError, ValueError):
                conf = -1

            has_arabic = bool(re.search(r"[\u0600-\u06FF]", token))
            if conf < 20 and len(token) <= 4 and not has_arabic:
                continue

            block = data.get("block_num", [0] * n)[i]
            par = data.get("par_num", [0] * n)[i]
            line = data.get("line_num", [0] * n)[i]
            key = (block, par, line)
            lines.setdefault(key, []).append(token)
        ordered = [" ".join(tokens) for _, tokens in sorted(lines.items()) if tokens]
        return "\n".join(ordered)

    def _ocr_with_confidence(self, image, config: str) -> Tuple[str, float, int, int]:
        try:
            data = pytesseract.image_to_data( image, lang=self.OCR_LANG, config=config, output_type=pytesseract.Output.DICT, )
            mean_conf, low_count, token_count = self._confidence_stats(data)
            cleaned = self._clean_low_confidence_tokens(data)
            cleaned = self._remove_ocr_layout_noise(self._normalize_ocr_text(cleaned))
            return cleaned, mean_conf, low_count, token_count
        except Exception as e:
            print(f"Tesseract confidence OCR error: {e}")
            return "", 0.0, 0, 0

    def _ocr_page(self, image) -> Tuple[str, float, int, int]:
        processed = self._preprocess_ocr_image(image)
        return self._ocr_with_confidence(processed, self.OCR_CONFIG)

    @staticmethod
    def _looks_like_heading(line: str) -> bool:
        line = line.strip()
        if not line or len(line) > 180:
            return False

        if PDFParser._PROCEDURE_HEADING_RE.match(line):
            return True
        if PDFParser._NUMBERED_HEADING_RE.match(line):
            return True
        if PDFParser._SECTION_HINT_RE.match(line):
            return True
        
        arabic = re.findall(r"[\u0600-\u06FF]", line)
        if 3 <= len(arabic) and len(line) <= 90:
            if not re.search(r"[؛:,.،.!?؟]", line) and not re.search(r"\d{2,}", line):
                heading_words = ( "إدارة", "قسم", "سياسة", "دليل", "إجراءات", "إجراء",
                                  "تقرير", "النماذج", "التقارير", "المسؤوليات", "الهدف",
                                  "النطاق", "التعريف", "الملاحق", "المخزون" )
                if any(word in line for word in heading_words):
                    return True
        return False

    @classmethod
    def _split_page_into_elements( cls, text: str, page_number: int, metadata: dict, ) -> List[Element]:
        elements: List[Element] = []
        buffer: List[str] = []

        def flush_paragraph() -> None:
            if not buffer:
                return
            paragraph = "\n".join(buffer).strip()
            if paragraph:
                elements.append( Element( type="paragraph", text=paragraph, page=page_number, metadata=dict(metadata), ))
            buffer.clear()

        for raw_line in text.splitlines():
            line = raw_line.strip()
            if not line:
                flush_paragraph()
                continue

            if cls._looks_like_heading(line):
                flush_paragraph()
                elements.append( Element( type="heading", text=line, level=2 if cls._PROCEDURE_HEADING_RE.match(line) else 3, page=page_number, metadata=dict(metadata), ))
            else:
                buffer.append(line)
        flush_paragraph()
        return elements

    @staticmethod
    def _table_area_ratio(page, tables) -> float:
        page_area = float(page.width * page.height)
        if page_area <= 0:
            return 0.0
        total = 0.0
        for table in tables:
            x0, top, x1, bottom = table.bbox
            total += max(0.0, x1 - x0) * max(0.0, bottom - top)
        return min(1.0, total / page_area)

    @staticmethod
    def _crop_bbox(image, bbox, page_width, page_height):
        x0, top, x1, bottom = bbox
        scale_x = image.width / float(page_width)
        scale_y = image.height / float(page_height)
        left = max(0, int(round(x0 * scale_x)))
        upper = max(0, int(round(top * scale_y)))
        right = min(image.width, int(round(x1 * scale_x)))
        lower = min(image.height, int(round(bottom * scale_y)))
        if right <= left or lower <= upper:
            return None
        return image.crop((left, upper, right, lower))

    def _ocr_table_cells(self, image, page, table) -> Optional[List[List[str]]]:
        rows = []
        cells = getattr(table, "cells", None)
        if not cells:
            return None

        row_groups = []
        for cell in cells:
            if not cell:
                continue
            x0, top, x1, bottom = cell
            placed = False
            for row in row_groups:
                if abs(row[0][1] - top) <= 4:
                    row.append(cell)
                    placed = True
                    break
            if not placed:
                row_groups.append([cell])

        row_groups.sort(key=lambda r: min(c[1] for c in r))
        for row_cells in row_groups:
            row_cells.sort(key=lambda c: c[0])
            row_text = []
            for cell_bbox in row_cells:
                crop = self._crop_bbox(image, cell_bbox, page.width, page.height)
                if crop is None:
                    row_text.append("")
                    continue

                if cv2 is not None and np is not None:
                    arr = np.array(crop)
                    if arr.ndim == 3:
                        arr = cv2.cvtColor(arr, cv2.COLOR_RGB2GRAY)
                    arr = cv2.copyMakeBorder(arr, 8, 8, 8, 8, cv2.BORDER_CONSTANT, value=255)
                    crop = arr

                cell_text, _conf, _low, _tokens = self._ocr_with_confidence( crop, self.CELL_OCR_CONFIG )
                row_text.append(cell_text.replace("\n", " ").strip())
            if any(row_text):
                rows.append(row_text)

        return rows or None

    @staticmethod
    def _mask_table_regions(image, page, tables):
        if not tables or cv2 is None or np is None:
            return image

        arr = np.array(image).copy()
        if arr.ndim == 2:
            fill = 255
        else:
            fill = (255, 255, 255)

        sx = image.width / float(page.width)
        sy = image.height / float(page.height)
        for table in tables:
            x0, top, x1, bottom = table.bbox
            left = max(0, int(round(x0 * sx)))
            upper = max(0, int(round(top * sy)))
            right = min(image.width, int(round(x1 * sx)))
            lower = min(image.height, int(round(bottom * sy)))
            arr[upper:lower, left:right] = fill

        return arr

    @staticmethod
    def _native_page_text(page) -> str:
        try:
            text = page.extract_text(x_tolerance=2, y_tolerance=3) or ""
            return PDFParser._normalize_ocr_text(text)
        except Exception:
            return ""

    @staticmethod
    def _page_ocr_is_weak(text: str, mean_conf: float, token_count: int) -> bool:
        if not text.strip():
            return True
        if token_count == 0:
            return True
        chars = re.sub(r"\s+", "", text)
        if len(chars) < 20:
            return True

        arabic_count = len(re.findall(r"[\u0600-\u06FF]", text))
        latin_count = len(re.findall(r"[A-Za-z]", text))
        if arabic_count == 0 and latin_count == 0:
            return True

        if mean_conf < 18 and arabic_count < 30:
            return True
        return False

    def parse(self, path: str) -> Document:
        import pdfplumber
        elements: List[Element] = []
        print(f"\nParsing PDF with OCR-first parser: {path}")

        with pdfplumber.open(path) as pdf:
            print(f"Pages: {len(pdf.pages)}")

            for page_number, page in enumerate(pdf.pages, start=1):
                try:
                    image = page.to_image(resolution=self.dpi).original
                except Exception as e:
                    print(f"Page {page_number}: OCR rendering failed: {e}")
                    continue
                try:
                    tables = page.find_tables()
                except Exception as e:
                    print(f"  Page {page_number}: table detection failed: {e}")
                    tables = []

                table_ratio = self._table_area_ratio(page, tables)
                metadata = {
                    "parser": "tesseract_ocr",
                    "ocr": True,
                    "ocr_model": self.OCR_LANG,
                    "ocr_dpi": self.dpi,
                    "ocr_psm": 6,
                    "ocr_confidence": True,
                    "table_detection": "pdfplumber_find_tables",
                    "table_area_ratio": round(table_ratio, 4),
                    "page": page_number,}

                page_image = self._mask_table_regions(image, page, tables)
                ocr_text, mean_conf, low_count, token_count = self._ocr_page(page_image)
                if self._page_ocr_is_weak(ocr_text, mean_conf, token_count):
                    native_text = self._native_page_text(page)
                    if native_text and len(native_text) > len(ocr_text):
                        print( f"  Page {page_number}: OCR weak " f"(conf={mean_conf:.1f}, tokens={token_count}); " "using native text fallback")
                        ocr_text = native_text
                        metadata["native_fallback"] = True
                    else:
                        metadata["native_fallback"] = False
                else:
                    metadata["native_fallback"] = False

                metadata["ocr_mean_confidence"] = round(mean_conf, 2)
                metadata["ocr_low_conf_short_tokens"] = low_count
                metadata["ocr_token_count"] = token_count

                if ocr_text:
                    page_elements = self._split_page_into_elements( ocr_text, page_number, metadata, )
                    elements.extend(page_elements)

                table_count = 0
                for table_index, table in enumerate(tables, start=1):
                    try:
                        rows = self._ocr_table_cells(image, page, table)
                    except Exception as e:
                        print( f"  Page {page_number}: table {table_index} OCR failed: {e}" )
                        rows = None

                    if rows:
                        table_metadata = dict(metadata)
                        table_metadata.update(
                            {  "element_type": "table",
                                "table_index": table_index,
                                "table_bbox": tuple(round(v, 2) for v in table.bbox), })
                        elements.append( Element( type="table", text="", table=Table(rows=rows), page=page_number, metadata=table_metadata, ))
                        table_count += 1

                print( f"  Page {page_number}: "
                       f"OCR conf={mean_conf:.1f}, tokens={token_count}, "
                       f"elements={len(elements)}, tables={table_count}")

        print(f"Finished OCR-first parsing: {path} -> {len(elements)} elements")
        return Document.create(source=path, elements=elements, format="pdf")

class DocxParser(BaseParser):
    def parse(self, path: str) -> Document:
        import docx
        from docx.oxml.ns import qn

        document = docx.Document(path)
        elements: List[Element] = []
        body = document.element.body

        for child in body.iterchildren():
            tag = child.tag
            if tag == qn("w:p"):
                para = next((p for p in document.paragraphs if p._p is child), None)
                if para is None or not para.text.strip():
                    continue
                style = (para.style.name or "").lower()
                match = re.match(r"heading (\d)", style)
                if match:
                    elements.append( Element( type="heading", text=para.text.strip(), level=int(match.group(1)), ))
                else:
                    elements.append(Element(type="paragraph", text=para.text.strip()))

            elif tag == qn("w:tbl"):
                table = next((t for t in document.tables if t._tbl is child), None)
                if table is None:
                    continue
                rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
                rows = [row for row in rows if any(row)]
                if rows:
                    elements.append(Element(type="table", text="", table=Table(rows=rows)))
        return Document.create(source=path, elements=elements, format="docx")

_EXTENSION_MAP = {
    ".txt": TextParser,
    ".md": MarkdownParser,
    ".markdown": MarkdownParser,
    ".html": HTMLParser,
    ".htm": HTMLParser,
    ".csv": CSVParser,
    ".pdf": PDFParser,
    ".docx": DocxParser,}

def get_parser_for(path: str) -> BaseParser:
    ext = os.path.splitext(path)[1].lower()
    try:
        parser_cls = _EXTENSION_MAP[ext]
    except KeyError as exc:
        raise ValueError( f"No parser registered for extension '{ext}' ({path})" ) from exc
    return parser_cls()

class DocumentParser:
    def parse(self, path: str) -> Document:
        return get_parser_for(path).parse(path)

    def parse_directory( self, dir_path: str, recursive: bool = True, ) -> List[Document]:
        documents: List[Document] = []

        print("\n" + "=" * 70)
        print(f"SCANNING DOCUMENT DIRECTORY: {dir_path}")
        print("=" * 70)

        if recursive:
            walker = os.walk(dir_path)
        else:
            walker = [(dir_path, [], os.listdir(dir_path))]
        file_count = 0
        for root, _dirs, files in walker:
            for name in sorted(files):
                ext = os.path.splitext(name)[1].lower()
                if ext not in _EXTENSION_MAP:
                    continue

                file_count += 1
                full_path = os.path.join(root, name)
                print(f"\n[{file_count}] Parsing: {full_path}")

                try:
                    document = self.parse(full_path)
                    documents.append(document)
                    print(f"    SUCCESS ({len(document.elements)} elements)")
                except Exception as e:
                    print(f"    ERROR: {e}")

        print("\n" + "=" * 70)
        print(f"DISCOVERED FILES: {file_count}")
        print(f"SUCCESSFULLY PARSED: {len(documents)}")
        print("=" * 70)
        return documents
